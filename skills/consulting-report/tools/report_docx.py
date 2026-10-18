# -*- coding: utf-8 -*-
"""보고서 docx 를 조립하는 함수 모음. 레이아웃 기본값이 들어 있다(rules/4-document.md).

  표지: 띠 없이 제목 한 줄(가운데, 20pt 굵게). 아래 오른쪽에 작성자와 작성일.
  글자: 본문, 표, 도표 제목, 소제목 12pt. 대제목 14pt. 출처와 주석 10pt.
  본문 문단은 양쪽 정렬, 줄 간격 1.35.
  그림은 본문 폭(170mm)에 맞추고 가운데 정렬. 도표 제목은 그림 위 왼쪽, 출처는 그림 바로 아래.
  블록 사이는 빈 줄(spacer)로 띄운다.
  표는 회색 머리줄 띠와 가로 실선만 쓴다.

쓰는 법 (보고서마다 빌드 스크립트 하나):
    import sys; sys.path.insert(0, "<이 폴더 경로>")
    import report_docx as rd
    doc = rd.new_document(title="...")
    rd.cover(doc, "제목", author="작성자", date="2026년 10월 25일")
    rd.h1(doc, "핵심 요약"); rd.body(doc, "문장")
    rd.save(doc, "보고서.docx")          # 이미 있는 파일은 덮어쓰지 않는다

글꼴은 기본 Freesentation 이다. 다른 글꼴을 쓰려면 환경변수 REPORT_FONT 에 글꼴 이름을 준다.
docx 는 글꼴 파일을 담지 않는다. 문서를 여는 컴퓨터에 그 글꼴이 설치돼 있어야 같은 모양으로 보인다.
"""
import copy
import os
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

FONT = os.environ.get('REPORT_FONT', 'Freesentation')
BODY_MM = 170                      # A4 에서 좌우 여백 20mm 를 뺀 본문 폭
CELL_MARGIN_DXA = 108

# 글자 크기(pt)
PT_TITLE, PT_H1, PT_H2, PT_BODY, PT_TABLE, PT_EXHIBIT, PT_SOURCE = 20, 14, 12, 12, 12, 12, 10

# 문서 뼈대 색
NAVY = RGBColor(0x1B, 0x3A, 0x6B)
BLUE = RGBColor(0x2E, 0x75, 0xB6)
GREY = RGBColor(0x55, 0x55, 0x55)
NAVY_HEX = '1B3A6B'
# 표와 도표 글자 색, 표 선 색
INK = RGBColor(0x2D, 0x2D, 0x2D)
NOTE = RGBColor(0x6B, 0x70, 0x7B)
TABLE_HEAD = 'DFE2E5'
TABLE_RULE = 'D5D8DC'

_ALIGN = {'left': WD_ALIGN_PARAGRAPH.LEFT, 'center': WD_ALIGN_PARAGRAPH.CENTER,
          'right': WD_ALIGN_PARAGRAPH.RIGHT, 'justify': WD_ALIGN_PARAGRAPH.JUSTIFY}


# ── 기본 서식 ─────────────────────────────────────────────────────────────────
def set_run(run, size=PT_BODY, bold=False, color=None, name=FONT):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rpr.append(rf)
    for a in ('w:eastAsia', 'w:ascii', 'w:hAnsi'):
        rf.set(qn(a), name)


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hexcolor)
    tcPr.append(shd)


def vcenter(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    va = OxmlElement('w:vAlign')
    va.set(qn('w:val'), 'center')
    tcPr.append(va)


def cell_fill(cell, text, size=PT_TABLE, bold=False, color=None, align='left', pad=3.5):
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = _ALIGN[align]
    p.paragraph_format.space_before = Pt(pad)
    p.paragraph_format.space_after = Pt(pad)
    set_run(p.add_run(text), size=size, bold=bold, color=color)
    vcenter(cell)


def fix_layout(t, widths_mm):
    """표를 본문 폭 안에 고정 폭으로 둔다(자동 맞춤 끔)."""
    t.autofit = False
    tbl = t._tbl
    tblPr = tbl.tblPr
    for tag in ('w:tblW', 'w:tblLayout', 'w:tblInd', 'w:tblCellMar'):
        el = tblPr.find(qn(tag))
        if el is not None:
            tblPr.remove(el)
    tblW = OxmlElement('w:tblW')
    tblW.set(qn('w:w'), str(int(sum(widths_mm) * 56.7))); tblW.set(qn('w:type'), 'dxa')
    tblPr.append(tblW)
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed')
    tblPr.append(lay)
    ind = OxmlElement('w:tblInd')
    ind.set(qn('w:w'), str(CELL_MARGIN_DXA)); ind.set(qn('w:type'), 'dxa')
    tblPr.append(ind)
    mar = OxmlElement('w:tblCellMar')
    for side in ('left', 'right'):
        el = OxmlElement(f'w:{side}')
        el.set(qn('w:w'), str(CELL_MARGIN_DXA)); el.set(qn('w:type'), 'dxa')
        mar.append(el)
    tblPr.append(mar)
    grid = tbl.find(qn('w:tblGrid'))
    if grid is not None:
        for gc, wmm in zip(grid.findall(qn('w:gridCol')), widths_mm):
            gc.set(qn('w:w'), str(int(wmm * 56.7)))
    for row in t.rows:
        for c, wmm in zip(row.cells, widths_mm):
            c.width = Mm(wmm)


# ── 문서와 표지 ────────────────────────────────────────────────────────────────
def new_document(title=None, author=None):
    """A4, 좌우 20mm, 위아래 18mm, 기본 글꼴 FONT."""
    doc = Document()
    st = doc.styles['Normal']
    st.font.name = FONT
    st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), FONT)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(20)
    sec.top_margin, sec.bottom_margin = Mm(18), Mm(18)
    doc.core_properties.author = author or ''
    if title:
        doc.core_properties.title = title
    return doc


def cover(doc, title, author=None, date=None):
    """표지. 제목은 가운데 20pt 굵게, 작성자와 작성일은 오른쪽 12pt 굵게."""
    p = doc.add_paragraph()
    p.alignment = _ALIGN['center']
    set_run(p.add_run(title), size=PT_TITLE, bold=True)
    for label, value in (('작성자', author), ('작성일', date)):
        if value:
            q = doc.add_paragraph()
            q.alignment = _ALIGN['right']
            set_run(q.add_run(f'{label}: {value}'), size=PT_BODY, bold=True)


# ── 제목과 본문 ────────────────────────────────────────────────────────────────
def h1(doc, text, page_break=False):
    """대제목. 14pt 남색 굵게, 아래 실선. page_break=True 면 새 쪽에서 시작."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    if page_break:
        p.paragraph_format.page_break_before = True
    set_run(p.add_run(text), size=PT_H1, bold=True, color=NAVY)
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    b = OxmlElement('w:bottom')
    b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), '14'); b.set(qn('w:space'), '4'); b.set(qn('w:color'), NAVY_HEX)
    pbdr.append(b); pPr.append(pbdr)
    return p


def h2(doc, text):
    """소제목. 12pt 파란색 굵게."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    set_run(p.add_run(text), size=PT_H2, bold=True, color=BLUE)
    return p


_QUOTED = re.compile(r'(‘[^’]*’)')


def body(doc, parts, keep=False, bold_quotes=False):
    """본문 문단. 12pt, 줄 간격 1.35, 양쪽 정렬.
    parts 는 문자열 또는 [(글, 굵게)] 목록. bold_quotes=True 면 작은따옴표로 묶은 콘텐츠 제목을 굵게 쓴다."""
    p = doc.add_paragraph()
    p.alignment = _ALIGN['justify']
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.35
    if keep:
        p.paragraph_format.keep_with_next = True
    if isinstance(parts, str):
        parts = ([(s, bool(_QUOTED.fullmatch(s))) for s in _QUOTED.split(parts) if s]
                 if bold_quotes else [(parts, False)])
    for text, bold in parts:
        set_run(p.add_run(text), size=PT_BODY, bold=bold)
    return p


def numbered(doc, items):
    """번호 목록. items 는 [(굵게 쓸 머리말, 나머지 글)] 또는 문자열 목록."""
    for i, item in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.35
        p.paragraph_format.left_indent = Mm(4)
        set_run(p.add_run(f'{i}. '), size=PT_BODY, bold=True)
        head, rest = item if isinstance(item, tuple) else ('', item)
        if head:
            set_run(p.add_run(head), size=PT_BODY, bold=True)
        set_run(p.add_run(rest), size=PT_BODY)


def spacer(doc):
    """빈 줄 하나. 블록 사이를 띄운다(넣는 자리는 rules/4-document.md)."""
    return doc.add_paragraph()


# ── 도표 ──────────────────────────────────────────────────────────────────────
def exhibit_title(doc, label, title, unit=None):
    """도표와 표 위 제목 줄. '도표 2. 제목 (단위: 회)' 모양(번호 보통, 제목 굵게, 단위 회색)."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    set_run(p.add_run(label + ' '), size=PT_EXHIBIT, color=INK)
    set_run(p.add_run(title), size=PT_EXHIBIT, bold=True, color=INK)
    if unit:
        set_run(p.add_run(' ' + unit), size=PT_EXHIBIT, color=NOTE)
    return p


def exhibit_img(doc, path, width_mm=BODY_MM, keep=True):
    """도표 그림. 기본은 본문 폭 가득, 가운데 정렬. keep=True 면 아래 출처 줄과 같은 쪽에 둔다."""
    p = doc.add_paragraph()
    p.alignment = _ALIGN['center']
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = keep
    p.add_run().add_picture(path, width=Mm(width_mm))
    return p


def source_line(doc, text, keep=False):
    """출처 줄. 10pt 회색, 그림 바로 아래. keep=True 면 아래 설명 문장과 같은 쪽에 둔다."""
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = keep
    set_run(p.add_run(text), size=PT_SOURCE, color=NOTE)
    return p


def note_line(doc, text):
    """표 아래 주석. '* ' 로 시작하는 10pt 회색 줄."""
    return source_line(doc, text if text.startswith('*') else '* ' + text)


def placeholder(doc, text):
    """나중에 사람이 그림을 넣을 자리."""
    t = doc.add_table(rows=1, cols=1)
    t.style = 'Table Grid'
    c = t.cell(0, 0)
    cell_fill(c, text, size=PT_BODY, bold=True, color=GREY, align='center')
    c.paragraphs[0].paragraph_format.space_before = Pt(28)
    c.paragraphs[0].paragraph_format.space_after = Pt(28)
    shade(c, 'F5F7FA')
    fix_layout(t, [BODY_MM])


# ── 표 ──────────────────────────────────────────────────────────────────────
def _borders(t):
    """세로줄 없음, 가로 실선만. 머리줄은 회색 띠가 구분한다."""
    tblPr = t._tbl.tblPr
    old = tblPr.find(qn('w:tblBorders'))
    if old is not None:
        tblPr.remove(old)
    b = OxmlElement('w:tblBorders')
    for side, val in (('top', 'nil'), ('left', 'nil'), ('bottom', 'single'), ('right', 'nil'),
                      ('insideH', 'single'), ('insideV', 'nil')):
        el = OxmlElement(f'w:{side}')
        el.set(qn('w:val'), val)
        if val == 'single':
            el.set(qn('w:sz'), '4'); el.set(qn('w:space'), '0'); el.set(qn('w:color'), TABLE_RULE)
        b.append(el)
    tblPr.append(b)


def data_table(doc, header, rows, widths, aligns, vmerge=(), bold_rows=(), note=None):
    """회색 머리줄 띠, 가로 실선, 12pt. 숫자 열은 aligns 에서 'right' 로 준다.
    widths 는 열 폭 비율(합이 본문 폭 170mm 로 맞춰진다).
    vmerge=[(열, 시작행, 끝행)] 은 세로로 합칠 칸(행 번호는 머리줄 뺀 0부터). bold_rows 는 굵게 쓸 행(합계).
    note 를 주면 표 아래 주석 줄을 쓰고, 표 마지막 줄과 같은 쪽에 묶는다."""
    scale = BODY_MM / sum(widths)
    widths = [w * scale for w in widths]
    t = doc.add_table(rows=len(rows) + 1, cols=len(header))
    _borders(t)
    for j, htxt in enumerate(header):
        c = t.cell(0, j)
        cell_fill(c, htxt, bold=True, color=INK, align=aligns[j])
        shade(c, TABLE_HEAD)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell_fill(t.cell(i + 1, j), str(val), bold=(i in bold_rows), color=INK, align=aligns[j])
    fix_layout(t, widths)
    for col, r0, r1 in vmerge:
        merged = t.cell(r0 + 1, col).merge(t.cell(r1 + 1, col))
        cell_fill(merged, rows[r0][col], color=INK, align=aligns[col])
    for row in t.rows:               # 한 행이 두 쪽에 걸쳐 쪼개지지 않게
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    keep_rows = t.rows if note else t.rows[:-1]
    for row in keep_rows:            # 표가 쪽 경계에서 갈라지지 않게 다음 행(주석이 있으면 주석)과 묶는다
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.keep_with_next = True
    if note:
        note_line(doc, note)
    return t


# ── 저장 ──────────────────────────────────────────────────────────────────────
def save(doc, path, overwrite=False):
    """저장. 이미 있는 파일은 덮어쓰지 않는다(Word 에서 사람이 고친 내용을 지키기 위해)."""
    if os.path.exists(path) and not overwrite:
        raise FileExistsError(
            f"이미 있는 파일은 덮어쓰지 않습니다: {path}\n"
            "Word 에서 고친 내용이 있을 수 있습니다. 다른 이름으로 저장하거나 overwrite=True 를 주세요.")
    doc.save(path)
    return path


# ── 사람이 고친 정본에 덧붙이기 ──────────────────────────────────────
# 정본(사람이 Word 에서 고친 docx)은 다시 빌드하지 않는다. 정본을 열어 이미 있는 문단, 그림, 표 줄을
# 본떠 새 블록을 끼워 넣고 다른 이름으로 저장한다. 서식은 본뜬 블록의 것이 그대로 따라오므로
# 사람이 바꾼 글자 크기, 간격, 그림 폭이 유지된다. 쓰는 법은 rules/4-document.md.

_W14 = '{http://schemas.microsoft.com/office/word/2010/wordml}'


def text_of(el):
    return ''.join(t.text or '' for t in el.iter(qn('w:t')))


def find_par(doc, startswith):
    """본문에서 글이 startswith 로 시작하는 첫 문단(w:p 요소)."""
    for p in doc.element.body.iterchildren(qn('w:p')):
        if text_of(p).startswith(startswith):
            return p
    raise LookupError(f'문단을 찾지 못했습니다: {startswith}')


def table_after(p):
    """문단 p 뒤에 나오는 첫 표(w:tbl 요소). 표 제목 문단으로 표를 찾을 때 쓴다."""
    el = p.getnext()
    while el is not None and el.tag != qn('w:tbl'):
        el = el.getnext()
    if el is None:
        raise LookupError(f'표를 찾지 못했습니다: {text_of(p)}')
    return el


def _scrub(el):
    """복제본에서 문단 고유 번호, 맞춤법 표시, 지난 쪽 나눔 표시를 지운다."""
    for e in el.iter():
        for a in (_W14 + 'paraId', _W14 + 'textId'):
            e.attrib.pop(a, None)
    for tag in ('w:proofErr', 'w:lastRenderedPageBreak'):
        for e in list(el.iter(qn(tag))):
            e.getparent().remove(e)
    return el


def clone_after(anchor, template):
    """template(문단이나 표)을 복제해 anchor 바로 뒤에 넣고 복제본을 돌려준다."""
    new = _scrub(copy.deepcopy(template))
    anchor.addnext(new)
    return new


def _is_bold(r):
    rpr = r.find(qn('w:rPr'))
    b = rpr.find(qn('w:b')) if rpr is not None else None
    return b is not None and b.get(qn('w:val')) not in ('0', 'false')


def _run_like(template, text, bold):
    r = copy.deepcopy(template)
    for child in list(r):
        if child.tag != qn('w:rPr'):
            r.remove(child)
    rpr = r.find(qn('w:rPr'))
    if rpr is None:
        rpr = OxmlElement('w:rPr')
        r.insert(0, rpr)
    for tag in ('w:b', 'w:bCs'):
        for e in rpr.findall(qn(tag)):
            rpr.remove(e)
    if bold:
        idx = sum(1 for c in rpr if c.tag in (qn('w:rStyle'), qn('w:rFonts')))
        rpr.insert(idx, OxmlElement('w:b'))
        rpr.insert(idx + 1, OxmlElement('w:bCs'))
    t = OxmlElement('w:t')
    t.text = text
    t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    r.append(t)
    return r


def add_parts(p, parts, clear=False, template=None):
    """문단 글을 쓴다. parts = [(글, 굵게 여부), ...]. 글자 서식은 그 문단의 첫 글자를 본뜬다.
    clear=True 면 원래 글을 지우고 쓴다. template 에 글자 묶음(w:r)을 주면 그 서식을 본뜬다."""
    runs = [r for r in p.iterchildren(qn('w:r'))]
    if template is None and not runs:
        raise ValueError(f'본뜰 글자가 없는 문단입니다: {text_of(p)}')
    template = copy.deepcopy(template if template is not None else runs[0])
    if clear:
        for r in runs:
            p.remove(r)
        for e in list(p.iterchildren(qn('w:proofErr'))):
            p.remove(e)
    for text, bold in parts:
        p.append(_run_like(template, text, bold))
    return p


def replace_in(p, old, new):
    """문단 안 한 글자 묶음(run) 속의 old 를 new 로 바꾼다. 없으면 멈춘다."""
    for t in p.iter(qn('w:t')):
        if t.text and old in t.text:
            t.text = t.text.replace(old, new, 1)
            return p
    raise LookupError(f'"{old}" 를 찾지 못했습니다: {text_of(p)}')


def swap_image(doc, p, path, width_emu=None):
    """그림 문단 p 의 그림을 path 로 바꾼다. 폭은 원래 그림 폭(또는 width_emu), 높이는 새 그림 비율을 따른다."""
    rid, img = doc.part.get_or_add_image(path)
    inline = next(p.iter(qn('wp:inline')))
    ext = inline.find(qn('wp:extent'))
    cx = int(width_emu or ext.get('cx'))
    cy = round(cx * img.px_height / img.px_width)
    ext.set('cx', str(cx))
    ext.set('cy', str(cy))
    for e in inline.iter(qn('a:ext')):
        if e.get('cx') is not None:
            e.set('cx', str(cx))
            e.set('cy', str(cy))
    next(inline.iter(qn('a:blip'))).set(qn('r:embed'), rid)
    ids = [int(d.get('id')) for d in doc.element.body.iter(qn('wp:docPr')) if (d.get('id') or '').isdigit()]
    new_id = max(ids + [0]) + 1
    docpr = inline.find(qn('wp:docPr'))
    docpr.set('id', str(new_id))
    docpr.set('name', f'Picture {new_id}')
    for c in inline.iter(qn('pic:cNvPr')):
        c.set('name', os.path.basename(path))
    return p


def append_row(tbl, values):
    """표 끝에 줄을 더한다. 마지막 줄을 본떠 칸마다 글만 바꾼다.
    원래 마지막 줄 문단에는 keepNext 를 붙여 다른 줄과 같게 맞춘다(마지막 줄만 keepNext 가 없다)."""
    last = tbl.findall(qn('w:tr'))[-1]
    new = _scrub(copy.deepcopy(last))
    for p in last.iter(qn('w:p')):
        ppr = p.find(qn('w:pPr'))
        if ppr is None:
            ppr = OxmlElement('w:pPr')
            p.insert(0, ppr)
        if ppr.find(qn('w:keepNext')) is None:
            idx = 1 if ppr.find(qn('w:pStyle')) is not None else 0
            ppr.insert(idx, OxmlElement('w:keepNext'))
    cells = new.findall(qn('w:tc'))
    if len(cells) != len(values):
        raise ValueError(f'칸 수가 다릅니다: 표 {len(cells)}칸, 값 {len(values)}개')
    for tc, v in zip(cells, values):
        add_parts(tc.find(qn('w:p')), [(v, False)], clear=True)
    last.addnext(new)
    return new


def set_run_texts(p, texts):
    """문단의 글자 묶음(w:r)마다 글을 바꾼다. 서식은 그대로 둔다. None 을 준 묶음은 지운다.
    예: 표 제목 `표 1. ` + 굵은 제목 + 회색 `(단위: 회)` 세 묶음을 그대로 본뜰 때."""
    runs = [r for r in p.iterchildren(qn('w:r')) if r.find(qn('w:t')) is not None]
    if len(runs) != len(texts):
        raise ValueError(f'글자 묶음 수가 다릅니다: 문단 {len(runs)}개, 값 {len(texts)}개 ({text_of(p)})')
    for r, text in zip(runs, texts):
        if text is None:
            p.remove(r)
            continue
        ts = r.findall(qn('w:t'))
        ts[0].text = text
        ts[0].set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
        for t in ts[1:]:
            r.remove(t)
    return p


def table_like(anchor, template, header, rows, widths=None):
    """template 표의 서식(머리줄 띠, 선, 칸 여백, 글자)을 본떠 열 수가 다른 새 표를 anchor 뒤에 만든다.
    widths 는 열 비율. 합계는 template 표 폭에 맞춘다. 마지막 줄만 keepNext 가 없는 것도 template 을 따른다."""
    t = clone_after(anchor, template)
    trs = t.findall(qn('w:tr'))
    if len(trs) < 3:
        raise ValueError('본뜰 표는 머리줄과 본문 두 줄 이상이어야 합니다')
    grid = t.find(qn('w:tblGrid'))
    total = sum(int(g.get(qn('w:w'))) for g in grid.findall(qn('w:gridCol')))
    n = len(header)
    ratios = widths or [1] * n
    ws = [int(total * w / sum(ratios)) for w in ratios]
    ws[-1] += total - sum(ws)
    for g in list(grid):
        grid.remove(g)
    for w in ws:
        g = OxmlElement('w:gridCol')
        g.set(qn('w:w'), str(w))
        grid.append(g)

    def row_like(tr_tpl, values, bold):
        tr = copy.deepcopy(tr_tpl)
        tc_tpl = tr.findall(qn('w:tc'))[0]
        for tc in tr.findall(qn('w:tc')):
            tr.remove(tc)
        for v, w in zip(values, ws):
            tc = copy.deepcopy(tc_tpl)
            tc.find(qn('w:tcPr')).find(qn('w:tcW')).set(qn('w:w'), str(w))
            p = tc.find(qn('w:p'))
            for extra in tc.findall(qn('w:p'))[1:]:
                tc.remove(extra)
            add_parts(p, [(v, bold)], clear=True)
            tr.append(tc)
        return tr

    head_tpl, body_tpl, last_tpl = trs[0], trs[1], trs[-1]
    head_bold = _is_bold(next(head_tpl.iter(qn('w:r'))))
    for tr in trs:
        t.remove(tr)
    t.append(row_like(head_tpl, header, head_bold))
    for i, r in enumerate(rows):
        if len(r) != n:
            raise ValueError(f'{i + 1}번째 줄 칸 수가 머리줄과 다릅니다')
        t.append(row_like(last_tpl if i == len(rows) - 1 else body_tpl, r, False))
    return t
